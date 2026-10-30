import datetime
import io
import tempfile
from pathlib import Path

from django.contrib.auth import get_user_model
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from . import services
from .models import (
    Assignment, Bank, Category, ContractType, Department, Direction, Employee, Entity, Pole, Position, Region,
    TempAgency,
)


class InitialValuesTests(TestCase):
    """The data migrations filled the lists."""

    def test_poles_and_their_entities(self):
        self.assertEqual(
            sorted(Entity.objects.filter(pole__name="SBD").values_list("name", flat=True)),
            ["Challenge", "Silent Believers Distribution"],
        )
        self.assertEqual(
            sorted(Entity.objects.filter(pole__name="Riaya").values_list("name", flat=True)),
            ["Eureka", "Norsudex", "Riaya"],
        )

    def test_simple_lists(self):
        self.assertEqual(Assignment.objects.count(), 5)
        self.assertEqual(ContractType.objects.count(), 4)
        self.assertEqual(TempAgency.objects.count(), 3)
        self.assertEqual(Region.objects.count(), 12)
        self.assertTrue(Bank.objects.filter(name="Umnia Bank").exists())

    def test_only_three_categories(self):
        self.assertEqual(
            sorted(Category.objects.values_list("name", flat=True)),
            ["Cadre", "Employé", "Ouvrier"],
        )

    def test_organisation_depends_on_the_pole(self):
        challenge = Entity.objects.get(name="Challenge")  # SBD: distribution
        eureka = Entity.objects.get(name="Eureka")  # Riaya: industry
        for entity in (challenge, eureka):
            self.assertTrue(entity.directions.filter(name="Direction des Ressources Humaines").exists())
        self.assertTrue(challenge.directions.filter(name="Direction Commerciale").exists())
        self.assertFalse(eureka.directions.filter(name="Direction Commerciale").exists())
        self.assertTrue(Position.objects.filter(
            name="Opérateur de production", department__direction__entity=eureka).exists())


class UniqueNameTests(TestCase):
    def test_same_name_with_other_case_is_refused(self):
        with self.assertRaises(ValidationError):
            Assignment(name="field").full_clean()  # "Field" already exists

    def test_same_direction_name_allowed_in_two_entities(self):
        sbd = Entity.objects.get(name="Silent Believers Distribution")
        riaya = Entity.objects.get(name="Riaya")
        Direction.objects.create(name="IT", entity=sbd)
        Direction(name="it", entity=riaya).full_clean()  # OK: another entity
        with self.assertRaises(ValidationError):
            Direction(name="it", entity=sbd).full_clean()  # same entity: refused

    def test_same_position_name_allowed_in_two_departments(self):
        sales = Department.objects.get(name="Ventes", direction__entity__name="Challenge")
        warehouse = Department.objects.get(name="Entreposage", direction__entity__name="Challenge")
        Position(name="chef des ventes", department=warehouse).full_clean()  # OK: another service
        with self.assertRaises(ValidationError):
            Position(name="chef des ventes", department=sales).full_clean()  # same service: refused

    def test_inactive_value_still_exists(self):
        pole = Pole.objects.get(name="SBD")
        pole.is_active = False
        pole.save()
        self.assertTrue(Entity.objects.filter(pole=pole).exists())


class ExportListsTests(TestCase):
    def test_one_sheet_per_list(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "lists.xlsx"
            call_command("export_lists", output=str(path), stdout=io.StringIO())
            workbook = load_workbook(path)

        self.assertEqual(len(workbook.sheetnames), 14)
        positions = workbook["Postes"]
        self.assertEqual([cell.value for cell in positions[1]], ["Entité", "Direction", "Service", "Nom", "Actif"])
        self.assertEqual(positions.max_row - 1, Position.objects.count())


class EmployeeTestCase(TestCase):
    def setUp(self):
        # "Ventes" of Challenge (pôle SBD), from the data migration.
        self.sales = Department.objects.get(name="Ventes", direction__entity__name="Challenge")
        self.salesman = self.sales.positions.get(name="Commercial terrain")

    def fiche(self, cin, **fields):
        """A valid fiche, not saved and without account; `fields` replaces any value."""
        values = {
            "title": Employee.Title.MR, "last_name": "Alami", "first_name": "Ahmed",
            "birth_date": datetime.date(1990, 5, 1), "cin": cin, "city": "Casablanca",
            "address": "12 rue Atlas", "phone": "0600000000", "professional_email": "ahmed.alami@sbd.ma",
            "marital_status": Employee.MaritalStatus.SINGLE,
            "department": self.sales, "position": self.salesman,
            "assignment": Assignment.objects.get(name="Field"), "category": Category.objects.get(name="Employé"),
            "contract_type": ContractType.objects.get(name="CDI"), "region": Region.objects.get(name="Casablanca-Settat"),
            "work_location": "Casablanca", "hire_date": datetime.date(2026, 1, 5),
            **fields,
        }
        return Employee(**values)

    def make_employee(self, username, cin, **fields):
        """A valid employee saved with its own account."""
        employee = self.fiche(cin, user=get_user_model().objects.create_user(username), **fields)
        employee.full_clean()
        employee.save()
        return employee


class EmployeeTests(EmployeeTestCase):
    def test_u_number_is_generated_in_order_and_never_changes(self):
        first = self.make_employee("a.alami", "AB111111")
        second = self.make_employee("b.bennani", "AB222222")

        self.assertEqual(first.u_number, f"U{first.pk:06d}")  # e.g. U000001
        self.assertEqual(int(second.u_number[1:]), int(first.u_number[1:]) + 1)

        first.last_name = "El Alami"
        first.save()
        first.refresh_from_db()
        self.assertEqual(first.u_number, f"U{first.pk:06d}")

    def test_leave_balance_is_18_days_by_default(self):
        # Le congé annuel légal au Maroc, quand le RH ne change rien.
        employee = self.make_employee("a.alami", "AB111111")
        employee.refresh_from_db()
        self.assertEqual(employee.initial_leave_balance, 18)

    def test_cin_is_unique_whatever_the_case(self):
        self.make_employee("a.alami", "AB111111")
        with self.assertRaises(ValidationError) as error:
            self.make_employee("b.bennani", "ab111111")
        self.assertIn("cin", error.exception.message_dict)

    def test_position_of_another_department_is_refused(self):
        warehouse_job = Position.objects.get(name="Magasinier", department__direction__entity__name="Challenge")
        with self.assertRaises(ValidationError) as error:
            self.make_employee("a.alami", "AB111111", position=warehouse_job)
        self.assertIn("position", error.exception.message_dict)

    def test_temp_agency_only_for_interim(self):
        interim = ContractType.objects.get(name="Intérimaire")
        agency = TempAgency.objects.get(name="PRESTATIS")

        with self.assertRaises(ValidationError):  # interim without agency
            self.make_employee("a.alami", "AB111111", contract_type=interim)
        with self.assertRaises(ValidationError):  # CDI with an agency
            self.make_employee("b.bennani", "AB222222", temp_agency=agency)
        self.make_employee("c.chraibi", "AB333333", contract_type=interim, temp_agency=agency)  # OK

    def test_direction_entity_and_pole_come_from_the_department(self):
        employee = self.make_employee("a.alami", "AB111111")
        self.assertEqual(employee.direction.name, "Direction Commerciale")
        self.assertEqual(employee.entity.name, "Challenge")
        self.assertEqual(employee.pole.name, "SBD")

    def test_n_plus_2_and_n_plus_3(self):
        director = self.make_employee("d.director", "AB000003")
        chief = self.make_employee("c.chief", "AB000002", manager=director)
        boss = self.make_employee("b.boss", "AB000001", manager=chief)
        employee = self.make_employee("a.alami", "AB111111", manager=boss)

        self.assertEqual(employee.manager, boss)
        self.assertEqual(employee.n_plus_2, chief)
        self.assertEqual(employee.n_plus_3, director)
        self.assertIsNone(director.n_plus_2)

    def test_manager_loop_is_refused(self):
        boss = self.make_employee("b.boss", "AB000001")
        employee = self.make_employee("a.alami", "AB111111", manager=boss)
        boss.manager = employee  # A -> B -> A
        with self.assertRaises(ValidationError) as error:
            boss.full_clean()
        self.assertIn("manager", error.exception.message_dict)


class AccountCreationTests(EmployeeTestCase):
    """Step 3: the login account is created from the fiche."""

    def test_username_is_firstname_dot_lastname_without_accents_or_spaces(self):
        self.assertEqual(services.make_username("Ahmed", "Alami"), "ahmed.alami")
        self.assertEqual(services.make_username("Fatima Ezzahra", "El Alami"), "fatima-ezzahra.el-alami")
        self.assertEqual(services.make_username("Aïcha", "Aït Taleb"), "aicha.ait-taleb")

    def test_same_name_gets_a_number(self):
        services.create_with_account(self.fiche("AB111111"))
        services.create_with_account(self.fiche("AB222222"))
        third = services.create_with_account(self.fiche("AB333333"))

        usernames = get_user_model().objects.filter(username__startswith="ahmed.alami").values_list("username", flat=True)
        self.assertEqual(sorted(usernames), ["ahmed.alami", "ahmed.alami2", "ahmed.alami3"])
        self.assertEqual(third.user.username, "ahmed.alami3")

    def test_account_has_cin_as_temporary_password(self):
        employee = services.create_with_account(self.fiche("ab123456"))
        user = employee.user

        self.assertTrue(user.check_password("AB123456"))  # the CIN, stored in capitals
        self.assertTrue(user.must_change_password)
        self.assertEqual(user.role, get_user_model().Role.COLLABORATOR)
        self.assertEqual(user.email, "ahmed.alami@sbd.ma")
        self.assertEqual(employee.u_number, f"U{employee.pk:06d}")

    def test_wrong_fiche_creates_no_account(self):
        warehouse_job = Position.objects.get(name="Magasinier", department__direction__entity__name="Challenge")
        with self.assertRaises(ValidationError):
            services.create_with_account(self.fiche("AB111111", position=warehouse_job))
        self.assertFalse(get_user_model().objects.filter(username="ahmed.alami").exists())

    def test_new_password_must_not_contain_the_cin(self):
        user = services.create_with_account(self.fiche("AB123456")).user
        for password in ("AB123456", "ab123456", "Sbd-ab123456!"):
            with self.subTest(password=password), self.assertRaises(ValidationError):
                validate_password(password, user=user)
        validate_password("Sbd-Nouveau-2026!", user=user)  # OK

    def admin_add(self, role):
        """L'IT crée une fiche dans /admin, avec ce rôle pour le compte."""
        it = get_user_model().objects.create_superuser("it.admin", password="Strong-Pass-2026!")
        self.client.force_login(it)
        fiche = self.fiche("CD987654")
        data = {
            "role": role, "status": fiche.status, "title": fiche.title, "last_name": "Bennani", "first_name": "Sara",
            "birth_date": "1995-03-10", "cin": "cd987654", "marital_status": fiche.marital_status,
            "phone": fiche.phone, "professional_email": "sara.bennani@sbd.ma", "address": fiche.address,
            "city": fiche.city, "country": fiche.country, "department": self.sales.pk, "position": self.salesman.pk,
            "contract_type": fiche.contract_type.pk, "category": fiche.category.pk,
            "assignment": fiche.assignment.pk, "region": fiche.region.pk, "work_location": fiche.work_location,
            "hire_date": "2026-01-05", "initial_leave_balance": "0",
            # Le petit tableau des enfants (inline) : l'admin veut savoir combien de lignes il y a.
            "children-TOTAL_FORMS": "0", "children-INITIAL_FORMS": "0",
        }
        response = self.client.post(reverse("admin:employees_employee_add"), data)
        self.assertEqual(response.status_code, 302)  # saved, back to the list
        return Employee.objects.get(cin="CD987654")

    def test_admin_add_creates_the_account(self):
        employee = self.admin_add("COLLABORATOR")
        self.assertEqual(employee.user.username, "sara.bennani")
        self.assertTrue(employee.user.check_password("CD987654"))
        self.assertEqual(employee.user.role, "COLLABORATOR")

    def test_admin_add_gives_the_chosen_role(self):
        """L'IT crée le premier Admin RH depuis la fiche."""
        self.assertEqual(self.admin_add("HR_ADMIN").user.role, "HR_ADMIN")

    def test_admin_change_page_shows_the_role(self):
        employee = self.admin_add("HR")
        response = self.client.get(reverse("admin:employees_employee_change", args=[employee.pk]))
        self.assertContains(response, "changer dans le compte")
        # La modification s'enregistre sans champ "Rôle du compte".
        self.assertNotContains(response, 'name="role"')
